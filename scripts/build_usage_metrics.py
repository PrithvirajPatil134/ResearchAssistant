"""
Build usage metric outputs for the three files in
'/Users/propatil/Desktop/Other Projects/TPM_Related/Products Usage Metrics/Excel Analysis'.

Targets:
1. Tableau TotalViews_ExcluFinTech_20260420.csv
   - Filter to the 13 dashboards from the image
   - Produce Total Views by dashboard by month
   - Write to a NEW companion XLSX file with a "Total Views by Month" sheet
     (CSV can't contain sheets, so a parallel xlsx is created).

2. Tableau Unique Viewers_Crosstab_ExclFinTech_20260420.xlsx
   - Filter to the same 13 dashboards
   - Produce Unique Viewers by dashboard by month (count distinct 'Viewer')
   - Add/replace a sheet "Unique Users by Month (Filtered)" in the same file.

3. QuickSight UniqueUsers&TotalView_ExcludingFinTech_20260420.xlsx
   - Column A = Dashboard Name (already filtered)
   - Total Views = count(eventname) by dashboard by month
   - Unique Viewers = count(distinct login_name) by dashboard by month
   - Add two new sheets: "Total Views by Month" and "Unique Users by Month".
"""
import os
import shutil
import pandas as pd
from openpyxl import load_workbook

BASE = "/Users/propatil/Desktop/Other Projects/TPM_Related/Products Usage Metrics/Excel Analysis"

# Target dashboard list from the image (Column B order preserved)
TARGET_DASHBOARDS = [
    "Inventory Reporting",
    "DEV Model Input Inspector",
    "Quest - QxG Retro Analysis - MBR View",
    "Quest - Planning Cycle Audit Review",
    "Golem - DEV Waterfall Bridging Dashboard",
    "Sasquatch LTV",
    "Yeti BOM",
    "Yeti BOM Volume and Tier Reports v3",
    "Yeti BOM Snapshot Data",
    "Yeti BOM programs with negative incremental volume",
    "RL-Repair Cost Report_V2",
    "IBVLF_Physical DEV_Dashboard",
    "Yeti BOM Live Data",
]


def month_str(ts):
    ts = pd.to_datetime(ts, errors="coerce")
    if pd.isna(ts):
        return None
    return ts.strftime("%Y-%m")


def pivot_by_month(df, dash_col, month_col, value_col, agg="sum", target_list=None):
    """Return a pivot: rows=dashboards, columns=months sorted asc, values=agg."""
    g = df.groupby([dash_col, month_col])[value_col]
    if agg == "sum":
        s = g.sum()
    elif agg == "count":
        s = g.count()
    elif agg == "nunique":
        s = g.nunique()
    else:
        raise ValueError(agg)
    pv = s.unstack(month_col).fillna(0).astype(int)
    # sort month columns chronologically
    pv = pv[sorted(pv.columns)]
    if target_list is not None:
        pv = pv.reindex(target_list).fillna(0).astype(int)
    pv.index.name = "Dashboard"
    return pv


# -----------------------------------------------------------------------------
# 1. Tableau Total Views (CSV)
# -----------------------------------------------------------------------------
print("Processing Tableau Total Views CSV ...")
csv_path = os.path.join(BASE, "Tableau TotalViews_ExcluFinTech_20260420.csv")
tv = pd.read_csv(csv_path, encoding="utf-16", sep="\t")
# Column2 = dashboard; Date Rollup = week start; '# of Views' = views
tv["_month"] = pd.to_datetime(tv["Date Rollup"], errors="coerce").dt.strftime("%Y-%m")
tv_filtered = tv[tv["Column2"].isin(TARGET_DASHBOARDS)].copy()
tv_pivot = pivot_by_month(
    tv_filtered,
    dash_col="Column2",
    month_col="_month",
    value_col="# of Views",
    agg="sum",
    target_list=TARGET_DASHBOARDS,
)
print(f"  Found {sum(tv_pivot.sum(axis=1) > 0)} / {len(TARGET_DASHBOARDS)} dashboards with data")

# Write to a companion XLSX next to the CSV
tv_out_path = os.path.join(BASE, "Tableau TotalViews_ExcluFinTech_20260420_ByMonth.xlsx")
with pd.ExcelWriter(tv_out_path, engine="openpyxl") as xw:
    tv_pivot.to_excel(xw, sheet_name="Total Views by Month")
print(f"  Wrote: {tv_out_path}")


# -----------------------------------------------------------------------------
# 2. Tableau Unique Viewers (XLSX) - add sheet
# -----------------------------------------------------------------------------
print("\nProcessing Tableau Unique Viewers XLSX ...")
uv_path = os.path.join(BASE, "Tableau Unique Viewers_Crosstab_ExclFinTech_20260420.xlsx")
# Real headers are on row 2 (index 1) per inspection
uv = pd.read_excel(uv_path, sheet_name="Sheet 1", header=1)
# Tableau crosstab uses merged cells: Project/Dashboard/Owner only populated on
# the first row of each group. Forward-fill so every per-viewer row carries them.
uv[["Project", "Dashboard", "Owner"]] = uv[["Project", "Dashboard", "Owner"]].ffill()
# Drop subtotal rows where Viewer/Department are both the literal string 'Total'
uv = uv[~((uv["Viewer"].astype(str) == "Total") & (uv["Department"].astype(str) == "Total"))]
# Identify weekly date columns (everything that isn't metadata)
meta_cols = {"Project", "Dashboard", "Owner", "Viewer", "Department", "Unnamed: 5"}
date_cols = [c for c in uv.columns if c not in meta_cols]
# Keep only rows with a Viewer id (per-user rows) and a target Dashboard
uv_rows = uv[uv["Dashboard"].isin(TARGET_DASHBOARDS) & uv["Viewer"].notna()].copy()

# Melt to long format: one row per (dashboard, viewer, week)
long = uv_rows.melt(
    id_vars=["Dashboard", "Viewer"],
    value_vars=date_cols,
    var_name="week",
    value_name="count",
)
long = long[long["count"].fillna(0) > 0].copy()
long["_month"] = pd.to_datetime(long["week"], errors="coerce").dt.strftime("%Y-%m")

uv_pivot = pivot_by_month(
    long,
    dash_col="Dashboard",
    month_col="_month",
    value_col="Viewer",
    agg="nunique",
    target_list=TARGET_DASHBOARDS,
)
print(f"  Found {sum(uv_pivot.sum(axis=1) > 0)} / {len(TARGET_DASHBOARDS)} dashboards with data")

# Write to same workbook, replacing if present
new_sheet = "Unique Users by Month (Filtered)"
wb = load_workbook(uv_path)
if new_sheet in wb.sheetnames:
    del wb[new_sheet]
wb.save(uv_path)
wb.close()
with pd.ExcelWriter(uv_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as xw:
    uv_pivot.to_excel(xw, sheet_name=new_sheet)
print(f"  Added sheet '{new_sheet}' to {os.path.basename(uv_path)}")


# -----------------------------------------------------------------------------
# 3. QuickSight XLSX - add two sheets
# -----------------------------------------------------------------------------
print("\nProcessing QuickSight XLSX ...")
qs_path = os.path.join(BASE, "QuickSight UniqueUsers&TotalView_ExcludingFinTech_20260420.xlsx")
# Real headers live on row 3 (index 2)
qs = pd.read_excel(qs_path, sheet_name="sheet1", header=2)
qs = qs[qs["Dashboard Name"].notna()].copy()
qs["_month"] = pd.to_datetime(qs["eventtime"], errors="coerce").dt.strftime("%Y-%m")
qs = qs[qs["_month"].notna()]

# Total Views = count of eventname per (dashboard, month)
qs_views = pivot_by_month(
    qs, dash_col="Dashboard Name", month_col="_month", value_col="eventname", agg="count"
)
# Unique Users = distinct login_name per (dashboard, month)
qs_users = pivot_by_month(
    qs, dash_col="Dashboard Name", month_col="_month", value_col="login_name", agg="nunique"
)
print(f"  Dashboards: {list(qs_views.index)}")
print(f"  Months: {list(qs_views.columns)}")

# Replace sheets if present
wb = load_workbook(qs_path)
for sname in ("Total Views by Month", "Unique Users by Month"):
    if sname in wb.sheetnames:
        del wb[sname]
wb.save(qs_path)
wb.close()
with pd.ExcelWriter(qs_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as xw:
    qs_views.to_excel(xw, sheet_name="Total Views by Month")
    qs_users.to_excel(xw, sheet_name="Unique Users by Month")
print(f"  Added sheets to {os.path.basename(qs_path)}")


# -----------------------------------------------------------------------------
# Summary
# -----------------------------------------------------------------------------
print("\n" + "=" * 72)
print("SUMMARY")
print("=" * 72)
print("\n[1] Tableau Total Views by Month (filtered)")
print(tv_pivot.to_string())
print("\n[2] Tableau Unique Viewers by Month (filtered)")
print(uv_pivot.to_string())
print("\n[3a] QuickSight Total Views by Month")
print(qs_views.to_string())
print("\n[3b] QuickSight Unique Users by Month")
print(qs_users.to_string())

# Flag dashboards without matches in the Tableau source
missing_tv = [d for d in TARGET_DASHBOARDS if tv_pivot.loc[d].sum() == 0]
missing_uv = [d for d in TARGET_DASHBOARDS if uv_pivot.loc[d].sum() == 0]
print("\nDashboards with zero Total Views (Tableau):", missing_tv)
print("Dashboards with zero Unique Viewers (Tableau):", missing_uv)
