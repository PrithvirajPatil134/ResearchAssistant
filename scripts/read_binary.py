#!/usr/bin/env python3
"""
Extract text from binary files (PDF, DOCX, XLSX).
Usage:
  python3 scripts/read_binary.py <path>                    # Full extraction
  python3 scripts/read_binary.py <path> --pages 1-5        # PDF pages 1-5
  python3 scripts/read_binary.py <path> --sheet Sheet1      # XLSX specific sheet
  python3 scripts/read_binary.py <path> --meta              # Metadata only (page count, structure)
"""
import sys
import argparse
import os


def extract_pdf(path, start_page=None, end_page=None, meta_only=False):
    import PyPDF2
    reader = PyPDF2.PdfReader(path)
    total = len(reader.pages)

    if meta_only:
        print(f"[PDF] {os.path.basename(path)}")
        print(f"Pages: {total}")
        return

    s = (start_page or 1) - 1
    e = min(end_page or total, total)
    print(f"[PDF: {total} pages, extracting {s + 1}-{e}]")

    for i in range(s, e):
        text = reader.pages[i].extract_text()
        if text:
            print(f"\n--- Page {i + 1} ---\n{text}")


def extract_docx(path, meta_only=False):
    import docx
    doc = docx.Document(path)
    total_paragraphs = len(doc.paragraphs)

    if meta_only:
        print(f"[DOCX] {os.path.basename(path)}")
        print(f"Paragraphs: {total_paragraphs}")
        sections = [p.text[:80] for p in doc.paragraphs if p.style.name.startswith("Heading")]
        if sections:
            print("Headings:")
            for h in sections:
                print(f"  - {h}")
        return

    print(f"[DOCX: {total_paragraphs} paragraphs]")
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            print(para.text)


def extract_xlsx(path, sheet_name=None, meta_only=False):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    if meta_only:
        print(f"[XLSX] {os.path.basename(path)}")
        print(f"Sheets: {wb.sheetnames}")
        for name in wb.sheetnames:
            ws = wb[name]
            print(f"  {name}: {ws.max_row} rows x {ws.max_column} cols")
        wb.close()
        return

    sheets = [sheet_name] if sheet_name else wb.sheetnames
    for name in sheets:
        if name not in wb.sheetnames:
            print(f"[WARNING] Sheet '{name}' not found. Available: {wb.sheetnames}")
            continue
        ws = wb[name]
        print(f"\n--- Sheet: {name} ({ws.max_row} rows x {ws.max_column} cols) ---")
        for row in ws.iter_rows(values_only=True):
            values = [str(v) if v is not None else "" for v in row]
            if any(v.strip() for v in values):
                print(" | ".join(values))
    wb.close()


def main():
    parser = argparse.ArgumentParser(description="Extract text from binary files")
    parser.add_argument("path", help="Path to file")
    parser.add_argument("--pages", help="Page range for PDF (e.g., 1-5)", default=None)
    parser.add_argument("--sheet", help="Sheet name for XLSX", default=None)
    parser.add_argument("--meta", action="store_true", help="Metadata only")
    args = parser.parse_args()

    path = args.path
    if not os.path.exists(path):
        print(f"[ERROR] File not found: {path}")
        sys.exit(1)

    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        start_page, end_page = None, None
        if args.pages:
            parts = args.pages.split("-")
            start_page = int(parts[0])
            end_page = int(parts[1]) if len(parts) > 1 else start_page
        extract_pdf(path, start_page, end_page, args.meta)

    elif ext == ".docx":
        extract_docx(path, args.meta)

    elif ext == ".xlsx":
        extract_xlsx(path, args.sheet, args.meta)

    else:
        print(f"[ERROR] Unsupported format: {ext}. Supported: .pdf, .docx, .xlsx")
        sys.exit(1)


if __name__ == "__main__":
    main()
